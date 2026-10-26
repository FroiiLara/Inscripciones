import os, sys, io, json, datetime as dt
os.environ.update(SECRET_KEY='k'*32, JWT_SECRET_KEY='j'*32, MONGO_URI='mongodb://x', FLASK_ENV='development', RATELIMIT_ENABLED='false')
import mongomock, pymongo
pymongo.MongoClient = mongomock.MongoClient
sys.path.insert(0, os.getcwd())
import app as A
from openpyxl import load_workbook
A.app.config['TESTING'] = True
sent = []
A.enviar_email = lambda to, subj, body: sent.append((to, subj, body))
A.cloudinary.uploader.upload = lambda f, **kw: {"secure_url": "https://res.cloudinary.com/x/y.pdf"}

c = A.app.test_client()
H = {}
with A.app.app_context():
    from flask_jwt_extended import create_access_token
    for mat, adm in (('1001', True), ('2002', False)):
        A.collection.insert_one(dict(matricula=mat, usuario='Admin Uno' if adm else 'Alumno Dos', email=f'{mat}@virtual.utsc.edu.mx', contrasena='x', es_admin=adm))
    tok = {m: {'Authorization': 'Bearer ' + create_access_token(identity=m)} for m in ('1001', '2002')}
today = dt.datetime.utcnow()
def fecha(delta_days): return (today - dt.timedelta(days=delta_days)).strftime("%Y-%m-%d %H:%M")
A.inscripciones_col.insert_many([
 dict(usuario='2002', folio='UTSC-AAA', nombre='=cmd|calc', curp='X', fecha_nacimiento='2000-05-01', telefono='811', carrera='ITID', estatus='En revisión', fecha_solicitud=fecha(3)),
 dict(usuario='3003', folio='UTSC-BBB', nombre='Beto', fecha_nacimiento='2004-01-01', carrera='II', estatus='Aprobada', fecha_solicitud=fecha(40)),
 dict(usuario='4004', folio='UTSC-CCC', nombre='Carla', fecha_nacimiento='2012-01-01', carrera='ITID', estatus='Rechazada', fecha_solicitud=fecha(70)),
])
A.reinscripciones_col.insert_many([
 dict(usuario='2002', folio='REINSC-AAA', matricula='2002', nombre='Alumno Dos', correo='a@b.c', carrera='Ingeniería Industrial', carrera_clave='II', cuatrimestre='4', es_estadia=False, materias=['Inglés IV','Ética'], estatus='En revisión', fecha_solicitud=fecha(2)),
 dict(usuario='3003', folio='REINSC-BBB', matricula='3003', nombre='Beto', carrera='Ing. en Mantenimiento Industrial', carrera_clave='IMI', cuatrimestre='11', es_estadia=True, materias=[], estatus='Aprobada', fecha_solicitud=fecha(10)),
])
ok = lambda cond, msg: print(('PASS ' if cond else 'FAIL ') + msg) or (cond or sys.exit(1))

# --- acceso ---
ok(c.get('/admin/dashboard', headers=tok['2002']).status_code == 403, 'dashboard bloqueado a alumnos')
ok(c.get('/admin/reportes/exportar', headers=tok['2002']).status_code == 403, 'exportar bloqueado a alumnos')
ok(c.get('/api/admin/dashboard', headers=tok['2002']).status_code == 403, 'api dashboard bloqueada a alumnos')

# --- páginas ---
r = c.get('/admin', headers=tok['1001']); ok(r.status_code == 200 and b'modalRechazo' in r.data and b'admin-subnav' in r.data, '/admin renderiza con modal y subnav')
r = c.get('/admin/dashboard', headers=tok['1001']); ok(r.status_code == 200 and b'chCarreras' in r.data, 'dashboard renderiza')
r = c.get('/admin/dashboard?estatus=Aprobada', headers=tok['1001']); ok(r.status_code == 200, 'dashboard con filtro')
r = c.get('/admin/reportes', headers=tok['1001']); ok(r.status_code == 200 and b'Descargar Excel' in r.data, 'reportes renderiza')

# --- stats API ---
d = c.get('/api/admin/dashboard', headers=tok['1001']).get_json()
ok(d['ok'] and d['totales'] == dict(inscripciones=3, reinscripciones=2, solicitudes=5, alumnos=3, estadias=1), f"totales {d['totales']}")
top = {x['clave']: x['total'] for x in d['por_carrera']}
ok(top['ITID'] == 2 and top['II'] == 2 and top['IMI'] == 1 and d['por_carrera'][0]['total'] == 2, 'conteo por carrera (clave y nombre)')
ok(d['estatus']['total'] == {'En revisión': 2, 'Aprobada': 2, 'Rechazada': 1}, 'estatus total')
ok(d['tasa_aprobacion'] == 66.7, f"tasa {d['tasa_aprobacion']}")
ok(d['edad_menor_18'] == 1 and sum(x['total'] for x in d['edades']) == 2, 'edades')
ok(sum(m['inscripciones'] for m in d['por_mes']) == 3 and len(d['por_mes']) == 12, 'por mes')
ok({q['cuatrimestre']: q['total'] for q in d['por_cuatrimestre']}[4] == 1 and len(d['por_cuatrimestre']) == 11, 'cuatrimestres')
d2 = c.get('/api/admin/dashboard?estatus=Aprobada', headers=tok['1001']).get_json()
ok(d2['totales']['solicitudes'] == 2, 'filtro aprobadas')
d3 = c.get('/api/admin/dashboard?estatus=Hack', headers=tok['1001']).get_json()
ok(d3['totales']['solicitudes'] == 5, 'filtro inválido => todas')

# --- Excel web ---
for tipo, hojas in (('general', ['Resumen','Por carrera','Inscripciones','Reinscripciones','Rechazadas']), ('inscripciones',['Inscripciones']), ('reinscripciones',['Reinscripciones']), ('por_carrera',['Por carrera']), ('rechazadas',['Rechazadas'])):
    r = c.get(f'/admin/reportes/exportar?tipo={tipo}', headers=tok['1001'])
    ok(r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet') and 'attachment' in r.headers['Content-Disposition'], f'xlsx {tipo} responde')
    wb = load_workbook(io.BytesIO(r.data)); ok(wb.sheetnames == hojas, f'hojas {tipo}: {wb.sheetnames}')
wb = load_workbook(io.BytesIO(c.get('/admin/reportes/exportar?tipo=inscripciones', headers=tok['1001']).data))
ws = wb['Inscripciones']; vals = [ws.cell(row=r, column=3).value for r in range(2, 5)]
ok('=cmd|calc' in vals and all(ws.cell(row=r, column=3).data_type == 's' for r in range(2, 5)), f'inyección de fórmulas neutralizada {vals}')
r = c.get('/admin/reportes/exportar?tipo=inscripciones&carrera=ITID&estatus=Rechazada', headers=tok['1001'])
ws = load_workbook(io.BytesIO(r.data))['Inscripciones']; ok(ws.max_row == 2 and ws['A2'].value == 'UTSC-CCC', 'filtros carrera+estatus en Excel')
r = c.get(f"/admin/reportes/exportar?tipo=inscripciones&desde={(today - dt.timedelta(days=5)).date()}", headers=tok['1001'])
ok(load_workbook(io.BytesIO(r.data))['Inscripciones'].max_row == 2, 'filtro desde')
r = c.get('/admin/reportes/exportar?tipo=zzz', headers=tok['1001']); ok(r.status_code == 302, 'tipo inválido => redirect')
r = c.get('/admin/reportes/exportar?desde=2026-02-01&hasta=2026-01-01', headers=tok['1001']); ok(r.status_code == 302, 'rango inválido => redirect')
gen = load_workbook(io.BytesIO(c.get('/admin/reportes/exportar?tipo=general', headers=tok['1001']).data))
rs = gen['Resumen']; ok(rs['A1'].value.startswith('Universidad'), 'hoja resumen'); ok(len(rs._charts) == 1 and len(gen['Por carrera']._charts) == 1, 'gráficas en Excel')

# --- enlace móvil ---
r = c.post('/api/admin/reportes/enlace', json={'tipo': 'general'}, headers=tok['2002']); ok(r.status_code == 403, 'enlace bloqueado a alumnos')
r = c.post('/api/admin/reportes/enlace', json={'tipo': 'nope'}, headers=tok['1001']); ok(r.status_code == 400, 'enlace tipo inválido')
r = c.post('/api/admin/reportes/enlace', json={'tipo': 'rechazadas'}, headers={**tok['1001'], 'X-Forwarded-Proto': 'https'}); j = r.get_json()
ok(r.status_code == 200 and j['url'].startswith('https://') and j['inscripciones'] == 3, f"enlace {j['url'][:50]}")
path = j['url'].split('localhost', 1)[1]
r = c.get(path); ok(r.status_code == 200 and r.mimetype.endswith('sheet'), 'descarga con token (sin JWT)')
def nuevo_enlace():
    return c.post('/api/admin/reportes/enlace', json={'tipo': 'rechazadas'}, headers=tok['1001']).get_json()['url'].split('localhost', 1)[1]
path = nuevo_enlace()
ok(c.get('/descargar_reporte/abc').status_code == 403, 'token inválido => 403')
A.REPORTE_ENLACE_SEGUNDOS = -1
ok(c.get(path).status_code == 403, 'token vencido => 403'); A.REPORTE_ENLACE_SEGUNDOS = 600
A.collection.update_one({'matricula': '1001'}, {'$set': {'es_admin': False}}); ok(c.get(path).status_code == 403, 'token de ex-admin => 403')
A.collection.update_one({'matricula': '1001'}, {'$set': {'es_admin': True}})

# --- rechazo con motivo ---
post = lambda **kw: c.post('/admin/actualizar_estatus', data=kw, headers=tok['1001'])
post(folio='UTSC-AAA', tipo='inscripcion', estatus='Rechazada', motivo='  ')
ok(A.inscripciones_col.find_one({'folio': 'UTSC-AAA'})['estatus'] == 'En revisión', 'rechazo sin motivo no se aplica (web)')
post(folio='UTSC-AAA', tipo='inscripcion', estatus='Rechazada', motivo='El acta de nacimiento no es actualizada <b>x</b>')
doc = A.inscripciones_col.find_one({'folio': 'UTSC-AAA'})
ok(doc['estatus'] == 'Rechazada' and doc['motivo_rechazo'].startswith('El acta') and doc['revisado_por'] == '1001', 'rechazo con motivo guardado')
mail = sent[-1]; ok('El acta de nacimiento no es actualizada' in mail[2] and '&lt;b&gt;' in mail[2] and mail[0].startswith('2002'), 'correo con motivo (escapado)')
r = c.get('/mi_historial', headers=tok['2002']); ok(b'El acta de nacimiento no es actualizada' in r.data and b'&lt;b&gt;' in r.data, 'alumno ve el motivo en su historial (escapado)')
h = c.get('/api/historial', headers=tok['2002']).get_json(); ok(h['inscripciones'][0]['motivo_rechazo'].startswith('El acta'), 'API historial incluye motivo_rechazo')
r = c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Rechazada'), headers=tok['1001']); ok(r.status_code == 400, 'API: rechazo sin motivo => 400')
r = c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Rechazada', motivo='Comprobante ilegible'), headers=tok['1001']); ok(r.status_code == 200, 'API: rechazo con motivo')
ok(c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Rechazada', motivo='x'*501), headers=tok['1001']).status_code == 400, 'API: motivo > 500 => 400')
ok(c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Pagada'), headers=tok['1001']).status_code == 400, 'API: estatus inválido => 400')
ok(c.post('/api/admin/actualizar_estatus', json=dict(folio='NOPE', tipo='reinscripcion', estatus='Aprobada'), headers=tok['1001']).status_code == 404, 'API: folio inexistente => 404')
ok(c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Aprobada'), headers=tok['2002']).status_code == 403, 'API: alumno no puede cambiar estatus')
r = c.post('/api/admin/actualizar_estatus', json=dict(folio='REINSC-AAA', tipo='reinscripcion', estatus='Aprobada'), headers=tok['1001'])
ok('motivo_rechazo' not in A.reinscripciones_col.find_one({'folio': 'REINSC-AAA'}), 'al aprobar se borra el motivo previo')
r = c.get('/admin', headers=tok['1001']); ok(b'motivo-rechazo' in r.data and b'El acta de nacimiento' in r.data, 'admin ve el motivo en la tabla')

# --- fecha de nacimiento ---
hoy = dt.datetime.utcnow().date()
def con_edad(a, extra=0): 
    try: return hoy.replace(year=hoy.year - a) - dt.timedelta(days=extra)
    except ValueError: return hoy.replace(year=hoy.year - a, day=28) - dt.timedelta(days=extra)
casos = [(con_edad(18), True, '18 años hoy'), (con_edad(18, -1), False, 'cumple 18 mañana'), (con_edad(17), False, '17 años'), (hoy, False, 'nació hoy'),
         (hoy + dt.timedelta(days=3), False, 'futura'), (con_edad(101), False, '101 años'), (con_edad(60), True, '60 años')]
for f, esperado, nombre in casos:
    ok(A.validar_fecha_nacimiento(f.isoformat())[0] == esperado, f'validador: {nombre}')
for malo in ('', 'hola', '2020-13-45', '20250-01-01'):
    ok(A.validar_fecha_nacimiento(malo)[0] is False, f'validador rechaza {malo!r}')
form = lambda fn: dict(nombre='Juan Perez Lopez', curp='PEGJ990101HNLRZN01', fecha_nacimiento=fn, telefono='8112345678', carrera='II', contacto_emergencia_nombre='Mama', contacto_emergencia_tel='8100000000',
                       fotografia=(io.BytesIO(b'\x89PNG\r\n\x1a\nxx'), 'f.png'), doc_acta=(io.BytesIO(b'%PDF-1.4 x'), 'a.pdf'), doc_cert=(io.BytesIO(b'%PDF-1.4 x'), 'c.pdf'), comprobante=(io.BytesIO(b'%PDF-1.4 x'), 'p.pdf'))
n0 = A.inscripciones_col.count_documents({})
r = c.post('/inscripcion_submit', data=form(con_edad(17).isoformat()), headers=tok['2002'], content_type='multipart/form-data')
ok(r.status_code == 302 and A.inscripciones_col.count_documents({}) == n0, 'servidor rechaza menor de 18 (no guarda)')
r = c.get('/inscripcion', headers=tok['2002'])
ok(r.status_code in (200, 500), f'/inscripcion {r.status_code}')
r = c.post('/inscripcion_submit', data=form(con_edad(18).isoformat()), headers=tok['2002'], content_type='multipart/form-data')
ok(A.inscripciones_col.count_documents({}) == n0 + 1, 'servidor acepta 18 años cumplidos')
print('ALL PASSED')
