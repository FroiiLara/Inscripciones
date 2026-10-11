"""
reportes.py — Estadísticas del dashboard y exportación a Excel
Sistema de Inscripciones UTSC

Funciones puras (no tocan la base de datos): reciben las listas de
solicitudes ya leídas de MongoDB y devuelven estadísticas o un archivo .xlsx.
"""
from collections import Counter, OrderedDict
from datetime import datetime, date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ESTATUS_LISTA = ['En revisión', 'Aprobada', 'Rechazada']

TIPOS_REPORTE = OrderedDict([
    ('general',         'Reporte general (todas las hojas)'),
    ('inscripciones',   'Detalle de inscripciones'),
    ('reinscripciones', 'Detalle de reinscripciones'),
    ('por_carrera',     'Resumen por carrera'),
    ('rechazadas',      'Solicitudes rechazadas con motivo'),
])

COLOR_VERDE = '1A5F5C'
COLOR_GRIS = 'F2F2F2'


# ======================================================
# UTILIDADES DE FECHAS Y CARRERAS
# ======================================================

def parsear_fecha_solicitud(texto):
    """'2026-09-15 14:30' -> datetime (o None si no se puede interpretar)."""
    try:
        return datetime.strptime(str(texto).strip()[:16], "%Y-%m-%d %H:%M")
    except (ValueError, TypeError):
        return None


def parsear_fecha_iso(texto):
    """'2026-09-15' -> date (o None)."""
    try:
        return datetime.strptime(str(texto).strip(), "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


def calcular_edad(fecha_nac, hoy=None):
    """Edad cumplida a la fecha de hoy. fecha_nac: date."""
    hoy = hoy or date.today()
    return hoy.year - fecha_nac.year - ((hoy.month, hoy.day) < (fecha_nac.month, fecha_nac.day))


def clave_carrera(doc, carreras):
    """Devuelve la clave (ITID, II, ...) de una solicitud, sin importar si
    se guardó la clave o el nombre completo. None si no se reconoce."""
    clave = doc.get('carrera_clave')
    if clave in carreras:
        return clave
    valor = str(doc.get('carrera') or '').strip()
    if valor in carreras:
        return valor
    for k, nombre in carreras.items():
        if nombre == valor:
            return k
    return None


def nombre_carrera(doc, carreras):
    clave = clave_carrera(doc, carreras)
    if clave:
        return carreras[clave]
    return str(doc.get('carrera') or '—')


# ======================================================
# FILTROS
# ======================================================

def filtrar_solicitudes(docs, carreras, estatus=None, carrera=None, desde=None, hasta=None):
    """Filtra una lista de solicitudes.
    estatus: texto exacto o None/''/'Todas'.  carrera: clave o None.
    desde / hasta: date (inclusive) sobre fecha_solicitud."""
    resultado = []
    for d in docs:
        if estatus and estatus != 'Todas' and d.get('estatus', 'En revisión') != estatus:
            continue
        if carrera and carrera != 'Todas' and clave_carrera(d, carreras) != carrera:
            continue
        if desde or hasta:
            f = parsear_fecha_solicitud(d.get('fecha_solicitud'))
            if f is None:
                continue
            if desde and f.date() < desde:
                continue
            if hasta and f.date() > hasta:
                continue
        resultado.append(d)
    return resultado


# ======================================================
# ESTADÍSTICAS PARA EL DASHBOARD
# ======================================================

RANGOS_EDAD = [('18-19', 18, 19), ('20-22', 20, 22), ('23-25', 23, 25),
               ('26-30', 26, 30), ('31+', 31, 200)]


def _conteo_estatus(docs):
    c = Counter(d.get('estatus', 'En revisión') for d in docs)
    return {e: c.get(e, 0) for e in ESTATUS_LISTA}


def _ultimos_meses(n, hoy=None):
    hoy = hoy or date.today()
    y, m = hoy.year, hoy.month
    meses = []
    for _ in range(n):
        meses.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    return list(reversed(meses))


def construir_estadisticas(inscripciones, reinscripciones, carreras, hoy=None):
    """Agrega los datos que usan las gráficas del dashboard (web y app)."""
    hoy = hoy or date.today()

    # --- Totales ---
    alumnos = {d.get('usuario') for d in inscripciones + reinscripciones if d.get('usuario')}
    est_insc = _conteo_estatus(inscripciones)
    est_rein = _conteo_estatus(reinscripciones)
    est_total = {e: est_insc[e] + est_rein[e] for e in ESTATUS_LISTA}
    resueltas = est_total['Aprobada'] + est_total['Rechazada']
    tasa = round(est_total['Aprobada'] * 100 / resueltas, 1) if resueltas else None

    # --- Por carrera ---
    por_carrera = OrderedDict((k, {'clave': k, 'nombre': n, 'inscripciones': 0,
                                   'reinscripciones': 0, 'total': 0})
                              for k, n in carreras.items())
    sin_clasificar = {'clave': '?', 'nombre': 'Sin clasificar', 'inscripciones': 0,
                      'reinscripciones': 0, 'total': 0}
    for tipo, docs in (('inscripciones', inscripciones), ('reinscripciones', reinscripciones)):
        for d in docs:
            fila = por_carrera.get(clave_carrera(d, carreras), sin_clasificar)
            fila[tipo] += 1
            fila['total'] += 1
    filas_carrera = list(por_carrera.values())
    if sin_clasificar['total']:
        filas_carrera.append(sin_clasificar)
    filas_carrera.sort(key=lambda f: (-f['total'], f['nombre']))
    carrera_top = filas_carrera[0]['nombre'] if filas_carrera and filas_carrera[0]['total'] else None

    # --- Reinscripciones por cuatrimestre ---
    cuatris = Counter()
    for d in reinscripciones:
        try:
            cuatris[int(d.get('cuatrimestre'))] += 1
        except (TypeError, ValueError):
            pass
    max_cuatri = max(list(cuatris) + [6])
    por_cuatrimestre = [{'cuatrimestre': q, 'total': cuatris.get(q, 0)}
                        for q in range(1, max_cuatri + 1)]

    # --- Solicitudes por mes (últimos 12 meses) ---
    meses = _ultimos_meses(12, hoy)
    por_mes = OrderedDict((m, {'mes': m, 'inscripciones': 0, 'reinscripciones': 0}) for m in meses)
    for tipo, docs in (('inscripciones', inscripciones), ('reinscripciones', reinscripciones)):
        for d in docs:
            f = parsear_fecha_solicitud(d.get('fecha_solicitud'))
            if f:
                clave = f"{f.year:04d}-{f.month:02d}"
                if clave in por_mes:
                    por_mes[clave][tipo] += 1

    # --- Edades de los aspirantes (inscripciones) ---
    edades = OrderedDict((r[0], 0) for r in RANGOS_EDAD)
    menores = 0
    for d in inscripciones:
        fn = parsear_fecha_iso(d.get('fecha_nacimiento'))
        if not fn:
            continue
        e = calcular_edad(fn, hoy)
        if e < 18:
            menores += 1
            continue
        for etiqueta, lo, hi in RANGOS_EDAD:
            if lo <= e <= hi:
                edades[etiqueta] += 1
                break

    return {
        'generado': datetime.utcnow().strftime("%Y-%m-%d %H:%M"),   # UTC, igual que fecha_solicitud
        'totales': {
            'inscripciones': len(inscripciones),
            'reinscripciones': len(reinscripciones),
            'solicitudes': len(inscripciones) + len(reinscripciones),
            'alumnos': len(alumnos),
            'estadias': sum(1 for d in reinscripciones if d.get('es_estadia')),
        },
        'estatus': {'inscripciones': est_insc, 'reinscripciones': est_rein, 'total': est_total},
        'tasa_aprobacion': tasa,
        'carrera_top': carrera_top,
        'por_carrera': filas_carrera,
        'por_cuatrimestre': por_cuatrimestre,
        'por_mes': list(por_mes.values()),
        'edades': [{'rango': k, 'total': v} for k, v in edades.items()],
        'edad_menor_18': menores,
    }


# ======================================================
# EXCEL
# ======================================================

_BORDE = Border(*(Side(style='thin', color='D9D9D9'),) * 4)
_FUENTE_ENC = Font(bold=True, color='FFFFFF')
_RELLENO_ENC = PatternFill('solid', fgColor=COLOR_VERDE)


def _texto(valor):
    """Valor seguro para una celda de texto (evita que '=...' se vuelva fórmula)."""
    if valor is None:
        return ''
    return str(valor)


def _escribir_texto(ws, fila, col, valor):
    c = ws.cell(row=fila, column=col, value=_texto(valor))
    c.data_type = 's'          # fuerza texto: un nombre como "=1+1" no se evalúa
    return c


def _encabezado(ws, fila, columnas):
    for i, nombre in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i, value=nombre)
        c.font = _FUENTE_ENC
        c.fill = _RELLENO_ENC
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = _BORDE
    ws.row_dimensions[fila].height = 28


def _ajustar_anchos(ws, minimo=10, maximo=48):
    anchos = {}
    for fila in ws.iter_rows():
        for c in fila:
            if c.value is None:
                continue
            largo = max(len(linea) for linea in str(c.value).split('\n'))
            anchos[c.column] = max(anchos.get(c.column, 0), largo)
    for col, largo in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimo, min(maximo, largo + 2))


def _hoja_tabla(wb, titulo, columnas, filas, textos=()):
    """Crea una hoja con encabezado, filas, filtro y panel congelado.
    textos: índices (0-based) de columnas que deben tratarse siempre como texto."""
    ws = wb.create_sheet(titulo)
    _encabezado(ws, 1, columnas)
    for r, fila in enumerate(filas, start=2):
        for c, valor in enumerate(fila, start=1):
            if isinstance(valor, str) or valor is None or (c - 1) in textos:
                celda = _escribir_texto(ws, r, c, valor)
            else:
                celda = ws.cell(row=r, column=c, value=valor)
            celda.border = _BORDE
            celda.alignment = Alignment(vertical='top', wrap_text=isinstance(valor, str) and len(valor) > 40)
            if r % 2 == 1:
                celda.fill = PatternFill('solid', fgColor=COLOR_GRIS)
    ws.freeze_panes = 'A2'
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
    _ajustar_anchos(ws)
    return ws


def _hoja_resumen(wb, stats, filtros_txt):
    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = 'Universidad Tecnológica de Santa Catarina'
    ws['A1'].font = Font(bold=True, size=14, color=COLOR_VERDE)
    ws['A2'] = 'Reporte de inscripciones y reinscripciones'
    ws['A2'].font = Font(bold=True, size=12)
    ws['A3'] = f"Generado: {stats['generado']} UTC"
    ws['A4'] = f"Filtros: {filtros_txt}"

    _encabezado(ws, 6, ['Indicador', 'Valor'])
    indicadores = [
        ('Inscripciones', stats['totales']['inscripciones']),
        ('Reinscripciones', stats['totales']['reinscripciones']),
        ('Total de solicitudes', stats['totales']['solicitudes']),
        ('Alumnos distintos', stats['totales']['alumnos']),
        ('Reinscripciones a estadía', stats['totales']['estadias']),
        ('Tasa de aprobación (%)', stats['tasa_aprobacion'] if stats['tasa_aprobacion'] is not None else 'N/D'),
        ('Carrera con más solicitudes', stats['carrera_top'] or 'N/D'),
    ]
    for i, (k, v) in enumerate(indicadores, start=7):
        _escribir_texto(ws, i, 1, k).border = _BORDE
        c = ws.cell(row=i, column=2, value=v)
        c.border = _BORDE
        c.alignment = Alignment(horizontal='left')

    inicio = 7 + len(indicadores) + 2
    _encabezado(ws, inicio, ['Estatus', 'Inscripciones', 'Reinscripciones', 'Total'])
    for i, e in enumerate(ESTATUS_LISTA, start=inicio + 1):
        _escribir_texto(ws, i, 1, e).border = _BORDE
        ws.cell(row=i, column=2, value=stats['estatus']['inscripciones'][e]).border = _BORDE
        ws.cell(row=i, column=3, value=stats['estatus']['reinscripciones'][e]).border = _BORDE
        ws.cell(row=i, column=4, value=f"=B{i}+C{i}").border = _BORDE

    pie = PieChart()
    pie.title = 'Solicitudes por estatus'
    pie.add_data(Reference(ws, min_col=4, min_row=inicio, max_row=inicio + 3), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=inicio + 1, max_row=inicio + 3))
    pie.height, pie.width = 7.5, 11
    ws.add_chart(pie, f"F{6}")

    _ajustar_anchos(ws, minimo=14, maximo=60)
    ws.column_dimensions['A'].width = 34
    return ws


def _hoja_por_carrera(wb, inscripciones, reinscripciones, carreras):
    columnas = ['Carrera', 'Inscripciones', 'Reinscripciones', 'Total',
                'Aprobadas', 'En revisión', 'Rechazadas']
    filas = []
    todas = [(d, 'i') for d in inscripciones] + [(d, 'r') for d in reinscripciones]
    claves = list(carreras.keys()) + ['?']
    for clave in claves:
        nombre = carreras.get(clave, 'Sin clasificar')
        propias = [(d, t) for d, t in todas if (clave_carrera(d, carreras) or '?') == clave]
        if clave == '?' and not propias:
            continue
        filas.append([
            nombre,
            sum(1 for _, t in propias if t == 'i'),
            sum(1 for _, t in propias if t == 'r'),
            None,
            sum(1 for d, _ in propias if d.get('estatus') == 'Aprobada'),
            sum(1 for d, _ in propias if d.get('estatus', 'En revisión') == 'En revisión'),
            sum(1 for d, _ in propias if d.get('estatus') == 'Rechazada'),
        ])
    ws = _hoja_tabla(wb, 'Por carrera', columnas, filas)
    n = len(filas)
    for r in range(2, n + 2):
        ws.cell(row=r, column=4, value=f"=B{r}+C{r}").border = _BORDE
    fila_tot = n + 2
    _escribir_texto(ws, fila_tot, 1, 'TOTAL').font = Font(bold=True)
    for col in range(2, 8):
        letra = get_column_letter(col)
        c = ws.cell(row=fila_tot, column=col, value=f"=SUM({letra}2:{letra}{n + 1})")
        c.font = Font(bold=True)
        c.border = _BORDE
    if n:
        ch = BarChart()
        ch.type = 'bar'
        ch.title = 'Solicitudes por carrera'
        ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ch.grouping, ch.overlap = 'stacked', 100
        ch.height, ch.width = 9, 20
        ws.add_chart(ch, f"A{fila_tot + 3}")
    ws.column_dimensions['A'].width = 58
    return ws


def _hoja_inscripciones(wb, inscripciones, carreras, hoy=None):
    columnas = ['Folio', 'Matrícula', 'Nombre', 'CURP', 'Fecha de nacimiento', 'Edad',
                'Teléfono', 'Carrera', 'Contacto de emergencia', 'Tel. contacto',
                'Fecha de solicitud', 'Estatus', 'Motivo de rechazo']
    filas = []
    for d in inscripciones:
        fn = parsear_fecha_iso(d.get('fecha_nacimiento'))
        filas.append([
            d.get('folio'), d.get('usuario'), d.get('nombre'), d.get('curp'),
            d.get('fecha_nacimiento'), calcular_edad(fn, hoy) if fn else None,
            d.get('telefono'), nombre_carrera(d, carreras),
            d.get('contacto_emergencia_nombre'), d.get('contacto_emergencia_tel'),
            d.get('fecha_solicitud'), d.get('estatus', 'En revisión'), d.get('motivo_rechazo'),
        ])
    return _hoja_tabla(wb, 'Inscripciones', columnas, filas, textos=(1, 6, 9))


def _hoja_reinscripciones(wb, reinscripciones, carreras):
    columnas = ['Folio', 'Matrícula', 'Nombre', 'Correo', 'Teléfono', 'Carrera',
                'Cuatrimestre', 'Estadía', 'Materias', 'Referencia de pago', 'Fecha de pago',
                'Fecha de solicitud', 'Estatus', 'Motivo de rechazo']
    filas = []
    for d in reinscripciones:
        try:
            cuatri = int(d.get('cuatrimestre'))
        except (TypeError, ValueError):
            cuatri = d.get('cuatrimestre')
        filas.append([
            d.get('folio'), d.get('matricula') or d.get('usuario'), d.get('nombre'),
            d.get('correo'), d.get('telefono'), nombre_carrera(d, carreras), cuatri,
            'Sí' if d.get('es_estadia') else 'No', ', '.join(d.get('materias') or []),
            d.get('referencia_pago'), d.get('fecha_pago'), d.get('fecha_solicitud'),
            d.get('estatus', 'En revisión'), d.get('motivo_rechazo'),
        ])
    return _hoja_tabla(wb, 'Reinscripciones', columnas, filas, textos=(1, 4, 9))


def _hoja_rechazadas(wb, inscripciones, reinscripciones, carreras):
    columnas = ['Tipo', 'Folio', 'Matrícula', 'Nombre', 'Carrera',
                'Fecha de solicitud', 'Fecha de revisión', 'Motivo de rechazo']
    filas = []
    for tipo, docs in (('Inscripción', inscripciones), ('Reinscripción', reinscripciones)):
        for d in docs:
            if d.get('estatus') != 'Rechazada':
                continue
            filas.append([tipo, d.get('folio'), d.get('matricula') or d.get('usuario'),
                          d.get('nombre'), nombre_carrera(d, carreras),
                          d.get('fecha_solicitud'), d.get('fecha_revision'),
                          d.get('motivo_rechazo') or '(sin motivo registrado)'])
    filas.sort(key=lambda f: str(f[5]), reverse=True)
    return _hoja_tabla(wb, 'Rechazadas', columnas, filas, textos=(2,))


def descripcion_filtros(estatus=None, carrera=None, desde=None, hasta=None, carreras=None):
    partes = []
    if estatus and estatus != 'Todas':
        partes.append(f"estatus = {estatus}")
    if carrera and carrera != 'Todas':
        partes.append(f"carrera = {(carreras or {}).get(carrera, carrera)}")
    if desde:
        partes.append(f"desde {desde.isoformat()}")
    if hasta:
        partes.append(f"hasta {hasta.isoformat()}")
    return '; '.join(partes) if partes else 'Sin filtros (todos los registros)'


def generar_excel(tipo, inscripciones, reinscripciones, carreras, filtros_txt='', hoy=None):
    """Genera el libro de Excel y lo devuelve como BytesIO listo para enviar.
    inscripciones / reinscripciones deben venir YA filtradas."""
    if tipo not in TIPOS_REPORTE:
        raise ValueError(f"Tipo de reporte desconocido: {tipo}")

    stats = construir_estadisticas(inscripciones, reinscripciones, carreras, hoy)
    wb = Workbook()

    if tipo == 'general':
        _hoja_resumen(wb, stats, filtros_txt)
        _hoja_por_carrera(wb, inscripciones, reinscripciones, carreras)
        _hoja_inscripciones(wb, inscripciones, carreras, hoy)
        _hoja_reinscripciones(wb, reinscripciones, carreras)
        _hoja_rechazadas(wb, inscripciones, reinscripciones, carreras)
    else:
        wb.remove(wb.active)
        if tipo == 'inscripciones':
            _hoja_inscripciones(wb, inscripciones, carreras, hoy)
        elif tipo == 'reinscripciones':
            _hoja_reinscripciones(wb, reinscripciones, carreras)
        elif tipo == 'por_carrera':
            _hoja_por_carrera(wb, inscripciones, reinscripciones, carreras)
        elif tipo == 'rechazadas':
            _hoja_rechazadas(wb, inscripciones, reinscripciones, carreras)

    wb.properties.title = TIPOS_REPORTE[tipo]
    wb.properties.creator = 'Sistema de Inscripciones UTSC'
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
